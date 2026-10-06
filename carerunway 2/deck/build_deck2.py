from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy

NAVY, LAVA, SLATE, MIST, LINE, TINT, WHITE = (RGBColor(0x11, 0x25, 0x30), RGBColor(0xFF, 0x36, 0x21), RGBColor(0x4E, 0x5E, 0x67),
                                              RGBColor(0x8A, 0x98, 0x9E), RGBColor(0xDD, 0xE3, 0xE6), RGBColor(0xF3, 0xF5, 0xF6), RGBColor(0xFF, 0xFF, 0xFF))
FONT = 'Arial'

prs = Presentation('tpl.pptx')
URL = 'https://claude.ai/artifact/KWxnv7oRJRx4aS9JTt8oXy'
URL_TEXT = 'claude.ai/artifact/KWxnv7oRJRx4aS9JTt8oXy'

def set_text(shape, text, size=None, bold=None, color=None):
    """Keep the first run's formatting, replace text, drop other paragraphs."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    r0 = p0.runs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    if size: r0.font.size = Pt(size)
    if bold is not None: r0.font.bold = bold
    if color is not None: r0.font.color.rgb = color

def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0):
    """paras: list of (text_or_runs, size, bold, color, space_after, align). runs = [(text, bold, color)]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Inches(margin))
    first = True
    for item in paras:
        text, size, bold, color = item[:4]
        after = item[4] if len(item) > 4 else 0
        align = item[5] if len(item) > 5 else PP_ALIGN.LEFT
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(after)
        runs = text if isinstance(text, list) else [(text, bold, color)]
        for t, b, c in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = c; r.font.name = FONT
    return box

def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    st = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: s._element.remove(st)
    if s.has_text_frame: s.text_frame.text = ''
    return s

def card(slide, x, y, w, h, off=0.09):
    """Signature motif: white card, navy outline, solid navy offset block behind."""
    rect(slide, x + off, y + off, w, h, fill=NAVY)
    return rect(slide, x, y, w, h, fill=WHITE, line=NAVY, lw=1.25)

def R(*parts):
    """parts: strings; wrap a string in [] (a list) to bold it."""
    return [((p[0], True, NAVY) if isinstance(p, list) else (p, False, NAVY)) for p in parts]

def qa(prompt, runs, size=13, after=10):
    """Template prompt as a quiet sentence-case label, answer in regular weight with one bold phrase."""
    return [(prompt, 10, False, MIST, 2), (runs, size, False, NAVY, after)]

def header(slide, n, section, claim, claim_size=28):
    shapes = {s.name: s for s in slide.shapes}
    set_text(shapes['TextBox 2'], f'SLIDE {n}   {section.upper()}', size=12)
    set_text(shapes['TextBox 3'], claim, size=claim_size)
    old = shapes['TextBox 4']; old._element.getparent().remove(old._element)

S1, S2, S3 = prs.slides
from pptx.oxml.ns import qn

# ---------------- Slide 1 ----------------
header(S1, 1, 'Problem & why it matters', "Seniors are growing faster than the clinics that serve them", claim_size=28)
left = []
left += qa('Team name', R(['[Team name]']))
left += qa('Problem statement', R('A3 HealthPulse: community healthcare demand forecasting'))
left += qa('The problem in one sentence', R('Some towns are running out of clinic capacity faster than others, and ',
           ['nobody can see which ones first,'], ' or whether the new polyclinics will open in time.'))
left += qa('Who is affected, and how badly', R("Singapore's 65+ population grew ", ['45% between 2015 and 2022'],
           ', to 735,000. Polyclinics and GP clinics grew 28% over the same years.'))
left += qa('Why it matters now', R('Healthier SG has signed up 1.4 million residents (May 2026), so more care lands on GPs. '
           'MOH is choosing polyclinic sites right now, and ', ["Bishan's has already slipped past 2030."]), after=0)
tb(S1, 0.9, 2.15, 6.1, 4.6, left)
cx, cy, cw, ch = 7.45, 2.2, 5.0, 3.15
card(S1, cx, cy, cw, ch)
tb(S1, cx + 0.3, cy + 0.2, cw - 0.6, 0.3, [('Seniors per clinic in Singapore', 11, True, SLATE)])
tb(S1, cx + 0.3, cy + 0.45, 2.3, 1.0, [('293', 60, True, LAVA)])
tb(S1, cx + 2.45, cy + 0.68, 2.3, 0.8, [('in 2022, up from', 12, False, SLATE), ('259 in 2015', 15, True, NAVY)])
bx, by, bw = cx + 0.3, cy + 1.75, cw - 0.6
tb(S1, bx, by, 1.6, 0.3, [('Seniors (65+)', 11, True, NAVY)])
rect(S1, bx + 1.55, by + 0.04, (bw - 2.35), 0.22, fill=LAVA)
tb(S1, bx + 1.55 + (bw - 2.35) + 0.08, by, 0.8, 0.3, [('+45%', 12, True, LAVA)])
tb(S1, bx, by + 0.45, 1.6, 0.3, [('Clinics', 11, True, NAVY)])
rect(S1, bx + 1.55, by + 0.49, (bw - 2.35) * 28 / 45, 0.22, fill=NAVY)
tb(S1, bx + 1.55 + (bw - 2.35) * 28 / 45 + 0.08, by + 0.45, 0.8, 0.3, [('+28%', 12, True, NAVY)])
tb(S1, bx, by + 0.95, bw, 0.3, [('Change from 2015 to 2022, from the DAISI data pack (population by age; MOH primary care facilities).', 8.5, False, MIST)])
sy = 5.6
tb(S1, cx, sy, cw, 0.3, [('And the pipeline is already slipping', 11, True, SLATE)])
rect(S1, cx, sy + 0.38, 2.25, 0.62, fill=TINT)
tb(S1, cx + 0.12, sy + 0.43, 2.05, 0.55, [('2025 plan', 9, False, SLATE), ('Bishan polyclinic by 2030', 11, True, NAVY)])
tb(S1, cx + 2.3, sy + 0.5, 0.45, 0.4, [('\u2192', 18, True, LAVA, 0, PP_ALIGN.CENTER)])
rect(S1, cx + 2.75, sy + 0.38, 2.25, 0.62, fill=RGBColor(0xFD, 0xE3, 0xDA))
tb(S1, cx + 2.87, sy + 0.43, 2.05, 0.55, [('2026 update', 9, False, SLATE), ('Pushed beyond 2030', 11, True, RGBColor(0xB9, 0x20, 0x0D))])

# ---------------- Slide 2 ----------------
header(S2, 2, 'Solution & data', 'CareRunway tells planners where the next clinic should go', claim_size=28)
left = []
left += qa('Your solution in one sentence', R('A planning tool that uses ', ['machine learning'],
           ' to forecast seniors per clinic in every area to 2031, checks MOH\'s announced polyclinics, and ranks where one new clinic helps most.'), size=11, after=7)
left += qa('Who uses it, and what decision it changes', R('Planners at the three regional health clusters, ',
           ['deciding where the next Healthier SG clinic goes.']), size=11, after=7)
left += [('Datasets we will use', 10, False, MIST, 2),
         (R('1. Residents by area, age and sex, 2000 to 2025 (SingStat)'), 10.5, False, NAVY, 2),
         (R('2. Clinic and polyclinic locations (data.gov.sg, MOH)'), 10.5, False, NAVY, 2),
         (R('3. Hospital admissions 1987 to 2026 (DAISI pack); admission rates by age (MOH)'), 10.5, False, NAVY, 4),
         ([('While cleaning the data, ', False, SLATE), ('we found 8 towns mislabelled', True, LAVA),
           (' in a government dataset. It overturned one of our own early findings.', False, SLATE)], 10, False, SLATE, 7)]
left += qa('Different from a generic dashboard', R('Most dashboards show where seniors live today. Ours says ',
           ['where the next clinic should go,'], ' checks MOH\'s own plans, and tests its forecast against the standard demographic method.'), size=11, after=9)
left += [([('Try the prototype: ', True, NAVY), (URL_TEXT, False, LAVA)], 10.5, False, NAVY, 1),
         ('Concept only: illustrative data, and many features are not built yet.', 9.5, False, MIST, 0)]
lbox = tb(S2, 0.9, 2.1, 4.5, 4.65, left)
for p in lbox.text_frame.paragraphs:
    for r in p.runs:
        if r.text == URL_TEXT: r.hyperlink.address = URL
img_h = 4.55; img_w = img_h * 2880 / 1952
rect(S2, 5.62, 2.12, img_w + 0.06, img_h + 0.06, fill=WHITE, line=LINE, lw=1)
S2.shapes.add_picture('mockup_v7_dashboard.png', Inches(5.65), Inches(2.15), Inches(img_w), Inches(img_h))

# ---------------- Slide 3 ----------------
header(S3, 3, 'Databricks architecture & impact', "Our data is already in Databricks. Here is how it becomes a recommendation.", claim_size=24)
steps = [('Sources', 'SingStat, data.gov.sg, MOH, DAISI pack, URA boundaries', False),
         ('Ingestion', 'Lakeflow Jobs land raw files in a Unity Catalog Volume', False),
         ('Transform', 'Delta bronze, silver, gold: label fixes, geocoding, catchments', False),
         ('Models', 'MLflow: ML area forecast, tested against the standard method; admissions forecast', True),
         ('Outputs', 'Site optimiser, AI/BI map dashboard, Genie space, App for briefs', False)]
x0, y0, w, h, gap = 0.9, 2.25, 2.12, 1.12, 0.27
for i, (t, desc, hero) in enumerate(steps):
    x = x0 + i * (w + gap)
    if hero: card(S3, x, y0, w, h, off=0.07)
    else: rect(S3, x, y0, w, h, fill=TINT)
    tb(S3, x + 0.14, y0 + 0.1, w - 0.28, h - 0.15, [(t, 12.5, True, LAVA if hero else NAVY, 3), (desc, 9.5, False, SLATE)])
    if i < len(steps) - 1:
        tb(S3, x + w - 0.02, y0 + h / 2 - 0.22, gap + 0.04, 0.4, [('\u2192', 16, True, MIST, 0, PP_ALIGN.CENTER)])
tb(S3, 0.9, y0 + h + 0.13, 11.6, 0.3, [(R("We use the ML forecast only where it wins: ", ['train to 2015, predict 2020; train to 2020, predict 2025,'],
     ' and compare its error with the standard method.'), 11, False, NAVY)])
# screenshot frame (team pastes a real Databricks screenshot here)
fx, fy, fw, fh = 0.9, 4.0, 5.5, 2.4
fr = rect(S3, fx, fy, fw, fh, fill=TINT, line=MIST, lw=1)
fr.line.dash_style = 4
tb(S3, fx + 0.3, fy + 0.85, fw - 0.6, 0.9, [('Paste a Databricks screenshot here', 13, True, SLATE, 3, PP_ALIGN.CENTER),
    ('Seniors per clinic, 2009 to 2022, queried from workspace.carerunway', 10, False, MIST, 0, PP_ALIGN.CENTER)])
right = []
right += qa("What we'll demo", R('A planner flips planned clinics on and off, clicks an area, asks Genie a question, and leaves with ',
            ['the top three sites and a one-page brief.']), size=11.5, after=7)
right += qa("How we'll measure impact", R(['Seniors living in under-served areas by 2031,'],
            ' with and without our recommended sites, and forecast error versus the standard method.'), size=11.5, after=7)
right += qa('What we can finish in two weeks', R('Week 1: pipeline, standard forecast, ', ['ML forecast and backtest.'],
            ' Week 2: site optimiser, planned-clinics check, dashboard, Genie. Stretch: nursing homes.'), size=11.5, after=0)
tb(S3, 6.75, 4.0, 5.75, 2.55, right)
tb(S3, 0.9, 6.55, 11.6, 0.25, [([('Team: ', True, NAVY), ('[Name] (data and models), [Name] (supply and mapping), [Name] (product and pitch)', False, SLATE)], 10, False, SLATE)])
S3.notes_slide.notes_text_frame.text = ("If the Databricks screenshot isn't ready, change the title to: "
    "'Built end to end on Databricks, with its error on screen' and delete the screenshot frame.")

prs.save('CareRunway_Round1.pptx')
print('saved')
