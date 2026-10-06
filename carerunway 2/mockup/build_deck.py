from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy

NAVY, LAVA, SLATE, MIST, LINE, TINT, WHITE = (RGBColor(0x11, 0x25, 0x30), RGBColor(0xFF, 0x36, 0x21), RGBColor(0x4E, 0x5E, 0x67),
                                              RGBColor(0x8A, 0x98, 0x9E), RGBColor(0xDD, 0xE3, 0xE6), RGBColor(0xF3, 0xF5, 0xF6), RGBColor(0xFF, 0xFF, 0xFF))
FONT = 'Arial'

prs = Presentation('tpl.pptx')

def set_text(shape, text, size=None, bold=None, color=None):
    """Keep the first run's formatting, replace text, drop other paragraphs."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    r0 = p0.runs[0]
    for r in p0.runs[1:]:
        r._r.getparent().remove(r._r)
    r0.text = text
    if size: r0.font.size = Pt(size)
    if bold is not None: r0.font.bold = bold
    if color is not None: r0.font.color.rgb = color

def tb(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, margin=0):
    """paras: list of (text_or_runs, size, bold, color, space_after, align). runs = [(text, bold, color)]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, side, Inches(margin))
    first = True
    for item in paras:
        text, size, bold, color = item[:4]
        after = item[4] if len(item) > 4 else 0
        align = item[5] if len(item) > 5 else PP_ALIGN.LEFT
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(after)
        runs = text if isinstance(text, list) else [(text, bold, color)]
        for t, b, c in runs:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.bold = b; r.font.color.rgb = c; r.font.name = FONT
    return box

def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None: s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    st = s._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if st is not None: s._element.remove(st)
    if s.has_text_frame: s.text_frame.text = ''
    return s

def card(slide, x, y, w, h, off=0.09):
    """Signature motif: white card, navy outline, solid navy offset block behind."""
    rect(slide, x + off, y + off, w, h, fill=NAVY)
    return rect(slide, x, y, w, h, fill=WHITE, line=NAVY, lw=1.25)

def qa(prompt, answer, size=13, after=9):
    """Template prompt (small grey caps) + loud answer."""
    return [(prompt.upper(), 9, True, MIST, 2), (answer, size, True, NAVY, after)]

def header(slide, n, section, claim, claim_size=28):
    shapes = {s.name: s for s in slide.shapes}
    set_text(shapes['TextBox 2'], f'SLIDE {n}   {section.upper()}', size=12)
    set_text(shapes['TextBox 3'], claim, size=claim_size)
    old = shapes['TextBox 4']; old._element.getparent().remove(old._element)

S1, S2, S3 = prs.slides

# ---------------- Slide 1 ----------------
header(S1, 1, 'Problem & why it matters', "Seniors are growing faster than the clinics that serve them", claim_size=28)
left = []
left += qa('Team name', '[Team name]')
left += qa('Problem statement', 'A3 HealthPulse: community healthcare demand forecasting')
left += qa('The problem in one sentence',
           'Seniors are outgrowing clinics unevenly across Singapore, and planners cannot see which areas run short first, '
           'or whether the clinics already announced will arrive in time.')
left += qa('Who is affected, and how badly',
           'From 2015 to 2022, the population aged 65+ grew 45% (506,000 to 735,000). '
           'Polyclinics and GP clinics grew only 28% (1,951 to 2,504).')
left += qa('Why it matters now',
           'Healthier SG has enrolled 1.4 million residents (May 2026), moving care onto GPs. MOH is choosing polyclinic sites now '
           '(28 today, 32 by 2030), and Bishan\'s has already slipped past 2030.', after=0)
tb(S1, 0.9, 2.15, 6.1, 4.6, [(p[0], p[1]+(0.5 if p[1]>12 else 0), *p[2:]) for p in left])

# hero card
cx, cy, cw, ch = 7.45, 2.2, 5.0, 3.15
card(S1, cx, cy, cw, ch)
tb(S1, cx + 0.3, cy + 0.2, cw - 0.6, 0.3, [('SENIORS PER CLINIC IN SINGAPORE', 9.5, True, SLATE)])
tb(S1, cx + 0.3, cy + 0.45, 2.3, 1.0, [('293', 60, True, LAVA)])
tb(S1, cx + 2.45, cy + 0.68, 2.3, 0.8, [('in 2022, up from', 12, False, SLATE), ('259 in 2015', 15, True, NAVY)])
# comparison bars
bx, by, bw = cx + 0.3, cy + 1.75, cw - 0.6
tb(S1, bx, by, 1.6, 0.3, [('Seniors (65+)', 11, True, NAVY)])
rect(S1, bx + 1.55, by + 0.04, (bw - 2.35), 0.22, fill=LAVA)
tb(S1, bx + 1.55 + (bw - 2.35) + 0.08, by, 0.8, 0.3, [('+45%', 12, True, LAVA)])
tb(S1, bx, by + 0.45, 1.6, 0.3, [('Clinics', 11, True, NAVY)])
rect(S1, bx + 1.55, by + 0.49, (bw - 2.35) * 28 / 45, 0.22, fill=NAVY)
tb(S1, bx + 1.55 + (bw - 2.35) * 28 / 45 + 0.08, by + 0.45, 0.8, 0.3, [('+28%', 12, True, NAVY)])
tb(S1, bx, by + 0.95, bw, 0.3, [('Change 2015 to 2022. Source: DAISI data pack (population by age; MOH primary care facilities).', 8.5, False, MIST)])
# Bishan slip strip
sy = 5.6
tb(S1, cx, sy, cw, 0.3, [('THE PIPELINE IS ALREADY SLIPPING', 9.5, True, SLATE)])
rect(S1, cx, sy + 0.38, 2.25, 0.62, fill=TINT)
tb(S1, cx + 0.12, sy + 0.43, 2.05, 0.55, [('2025 plan', 9, False, SLATE), ('Bishan polyclinic by 2030', 11, True, NAVY)])
tb(S1, cx + 2.3, sy + 0.5, 0.45, 0.4, [('\u2192', 18, True, LAVA, 0, PP_ALIGN.CENTER)])
rect(S1, cx + 2.75, sy + 0.38, 2.25, 0.62, fill=RGBColor(0xFD, 0xE3, 0xDA))
tb(S1, cx + 2.87, sy + 0.43, 2.05, 0.55, [('2026 update', 9, False, SLATE), ('Pushed beyond 2030', 11, True, RGBColor(0xB9, 0x20, 0x0D))])

# ---------------- Slide 2 ----------------
header(S2, 2, 'Solution & data', 'CareRunway shows where seniors will run short of clinics, and where the next clinic does the most good', claim_size=24)
left = []
left += qa('Your solution in one sentence',
           'A planning tool that forecasts seniors per clinic in every planning area to 2031, grades MOH\'s planned clinics, '
           'and recommends where the next clinic helps most.', size=11.5, after=8)
left += qa('Who uses it, and what decision it changes',
           'Regional health cluster planners (NHG Health, NUHS, SingHealth) deciding where to open the next Healthier SG clinic, '
           'and MOH checking whether its pipeline lands in time.', size=11.5, after=8)
left += [('DATASETS (NAME + SOURCE)', 9, True, MIST, 2),
         ('1. Residents by planning area, age and sex, 2000 to 2025 (SingStat)', 11, True, NAVY, 2),
         ('2. CHAS clinic and polyclinic locations (data.gov.sg, MOH)', 11, True, NAVY, 2),
         ('3. Public hospital admissions 1987 to 2026 (DAISI pack) and admission rates by age (MOH)', 11, True, NAVY, 3),
         ([('Plus URA planning-area boundaries and HDB elderly by town, where ', False, SLATE), ('we found and fixed 8 mislabelled towns', True, LAVA),
           (' in the source.', False, SLATE)], 10, False, SLATE, 8)]
left += qa('Different from a generic dashboard',
           'It answers "when" and "where next", not "where now". It grades MOH\'s own pipeline. And it shows its own error and data dates.',
           size=11.5, after=0)
tb(S2, 0.9, 2.32, 4.8, 4.4, left)
img_w = 6.75; img_h = img_w * 500 / 900
card(S2, 5.95, 2.55, img_w + 0.1, img_h + 0.1, off=0.09)
S2.shapes.add_picture('mockup_v4_slide.png', Inches(6.0), Inches(2.6), Inches(img_w), Inches(img_h))
tb(S2, 5.95, 2.55 + img_h + 0.25, img_w, 0.3,
   [('Concept mockup with illustrative values. Planners switch planned clinics on or off, pick an area, and ask Genie follow-up questions.', 9, False, MIST)])

# ---------------- Slide 3 ----------------
header(S3, 3, 'Databricks architecture & impact', 'Built end to end on Databricks, with its error on screen', claim_size=28)
steps = [('Sources', 'SingStat, data.gov.sg, MOH, DAISI pack, URA boundaries', False),
         ('Ingestion', 'Lakeflow Jobs load raw files into a Unity Catalog Volume', False),
         ('Transform', 'Delta bronze, silver, gold: label fixes, geocoding, catchments', False),
         ('Models', 'MLflow: population forecast, admissions forecast, next-site optimiser', True),
         ('Outputs', 'AI/BI map dashboard, Genie space, App for planning briefs', False)]
x0, y0, w, h, gap = 0.9, 2.15, 2.12, 1.1, 0.27
for i, (t, desc, hero) in enumerate(steps):
    x = x0 + i * (w + gap)
    if hero:
        card(S3, x, y0, w, h, off=0.07)
    else:
        rect(S3, x, y0, w, h, fill=TINT)
    tb(S3, x + 0.14, y0 + 0.12, w - 0.28, h - 0.2, [(t, 13, True, LAVA if hero else NAVY, 4), (desc, 10, False, SLATE)])
    if i < len(steps) - 1:
        tb(S3, x + w - 0.02, y0 + h / 2 - 0.22, gap + 0.04, 0.4, [('\u2192', 16, True, MIST, 0, PP_ALIGN.CENTER)])
tb(S3, 0.9, y0 + h + 0.2, 11.6, 0.3,
   [([('Validation built in: ', True, NAVY), ('backtest on 2015 data, predicting 2025 by area; area totals must match actual national admissions. '
      'The error is shown on the dashboard.', False, SLATE)], 10.5, False, SLATE)])
cols = [('The user-facing output we will demo',
         'A planner picks an area, sees the year it becomes under-served, switches planned clinics on and off, asks Genie a follow-up, '
         'and gets the three best sites for the next clinic with a one-page brief.'),
        ('Measurable impact if it worked',
         'Seniors living in under-served areas by 2031, with and without the recommended sites. Years of warning before an area '
         'becomes under-served. Planned clinics flagged as late or ahead of need.'),
        ('What we can finish in the two-week sprint',
         'Week 1: pipeline, population forecast and backtest. Week 2: next-site optimiser, planned-clinics check, dashboard and Genie space. '
         'Stretch: nursing homes view, AI planning brief.')]
cw3, cy3, chh = 3.72, 4.1, 1.9
for i, (p, a) in enumerate(cols):
    x = 0.9 + i * (cw3 + 0.24)
    if i == 1:
        card(S3, x, cy3, cw3, chh, off=0.08)
    else:
        rect(S3, x, cy3, cw3, chh, fill=WHITE, line=LINE, lw=1)
    tb(S3, x + 0.2, cy3 + 0.16, cw3 - 0.4, chh - 0.25, [(p.upper(), 9, True, LAVA if i == 1 else MIST, 6), (a, 12.5, True, NAVY)])

prs.save('CareRunway_Round1.pptx')
print('saved')
